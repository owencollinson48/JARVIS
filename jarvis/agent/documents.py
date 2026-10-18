"""Conservative DOCX/XLSX edits and PDF overlays; preserve originals and verify saves."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from . import filesystem as fs, policy



def check_archive(raw):
    with ZipFile(BytesIO(raw)) as archive:
        sizes=[item.file_size for item in archive.infolist()]
        if len(sizes)>3000 or sum(sizes)>50_000_000 or any(size>20_000_000 for size in sizes):
            raise ValueError('This Office document expands beyond the supported parsing limit.')


def paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:yield from cell.paragraphs


def read(ctx,a):
    policy.require(ctx,1,'Read this approved document.')
    path,raw=fs.read_local(a['root'],a['path'])
    if policy.financial_path(path):policy.require(ctx,3,'Share this potentially financial document with OpenAI? Authentication/payment credentials remain withheld.')
    if not policy.settings()['share_files'] and not ctx.approved:
        raise policy.ApprovalRequired('Share the extracted text/cells from this document with OpenAI?')
    suffix=path.suffix.lower()
    if suffix in ('.docx','.xlsx'):check_archive(raw)
    if suffix=='.docx':
        from docx import Document
        doc=Document(BytesIO(raw));text='\n'.join(p.text for p in paragraphs(doc))
        text+='\nBODY PARAGRAPH STYLES: '+str([dict(index=i,style=p.style.name,text=p.text[:80]) for i,p in enumerate(doc.paragraphs[:100])])
    elif suffix=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(raw),read_only=True,data_only=False)
        try:
            lines=[]
            for sheet in book:
                for row in sheet.iter_rows(max_row=min(sheet.max_row,100),max_col=min(sheet.max_column,20)):
                    for cell in row:
                        if cell.value is not None:lines.append(f'{sheet.title}!{cell.coordinate}: {cell.value}')
            text='\n'.join(lines)
        finally:book.close()
    elif suffix=='.pdf':
        from pypdf import PdfReader
        pdf=PdfReader(BytesIO(raw))
        if pdf.is_encrypted:raise ValueError('Encrypted documents require manual authentication.')
        text='\n'.join((p.extract_text() or '') for p in pdf.pages[:30])
    else:return fs.read(ctx,a)
    fs.remember(ctx,a['root'],a['path'])
    return policy.guard_data(dict(text=text[:12000],truncated=len(text)>12000,version=fs.digest(raw),path=a['path'],scope='DOCX body/tables; XLSX first 100 rows/20 columns per sheet; PDF first 30 pages.'))


def word_replace(ctx,a):
    policy.computer_access(ctx);policy.require(ctx,2,'Replace one Word text span with an original backup.')
    path,raw=fs.read_local(a['root'],a['path'])
    if path.suffix.lower()!='.docx':raise ValueError('Word editing supports DOCX only.')
    check_archive(raw)
    from docx import Document
    doc=Document(BytesIO(raw));found=[p for p in paragraphs(doc) if a['old'] and a['old'] in p.text]
    if len(found)!=1 or found[0].text.count(a['old'])!=1:raise ValueError('Choose a unique exact text span in a body paragraph/table cell. No formatting was guessed.')
    paragraph=found[0]
    if ''.join(run.text for run in paragraph.runs)!=paragraph.text:raise ValueError('This paragraph contains unsupported hyperlinks/fields. Use Word directly; original retained.')
    expected=paragraph.text.replace(a['old'],a['new'],1)
    start=paragraph.text.index(a['old']);end=start+len(a['old']);position=0;inserted=False
    for run in paragraph.runs:
        text=run.text;left=position;right=position+len(text);position=right
        if right<=start or left>=end:continue
        prefix=text[:max(0,start-left)];suffix=text[max(0,end-left):] if right>end else ''
        run.text=prefix+(a['new'] if not inserted else '')+suffix;inserted=True
    if not inserted or paragraph.text!=expected:raise ValueError('Unsupported Word text structure; original retained.')
    output=BytesIO();doc.save(output)
    verified=Document(BytesIO(output.getvalue()))
    if not any(a['new'] in p.text for p in paragraphs(verified)):raise OSError('Word save verification failed.')
    return fs.save_bytes(ctx,a['root'],a['path'],output.getvalue(),a['version'],'docx')


def excel_update(ctx,a):
    policy.computer_access(ctx)
    if any(c['value'].startswith(('=','+','@')) for c in a['cells']):
        policy.require(ctx,3,'These cells contain formulas or external-link-capable content. Review the exact expressions.')
    policy.require(ctx,2,'Update these spreadsheet cells with an original backup.')
    path,raw=fs.read_local(a['root'],a['path'])
    if path.suffix.lower()!='.xlsx':raise ValueError('Spreadsheet editing supports ordinary XLSX, not macros/XLSM.')
    check_archive(raw)
    from openpyxl import load_workbook
    book=load_workbook(BytesIO(raw));sheet=book[a['sheet']]
    for item in a['cells']:
        import re
        from openpyxl.utils.cell import coordinate_to_tuple
        if not re.fullmatch(r'[A-Za-z]{1,3}[1-9][0-9]{0,6}',item['cell']):raise ValueError('Use a single A1-style cell address.')
        row,column=coordinate_to_tuple(item['cell'])
        if row>1048576 or column>16384:raise ValueError('Cell is outside the Excel sheet bounds.')
        value=item['value']
        if item['type']=='number':value=float(value)
        elif item['type']=='blank':value=None
        sheet[item['cell']]=value
    output=BytesIO();book.save(output);book.close()
    check=load_workbook(BytesIO(output.getvalue()))
    try:
        for item in a['cells']:
            expected=float(item['value']) if item['type']=='number' else None if item['type']=='blank' else item['value']
            if check[a['sheet']][item['cell']].value!=expected:raise OSError('Spreadsheet verification failed.')
    finally:check.close()
    result=fs.save_bytes(ctx,a['root'],a['path'],output.getvalue(),a['version'],'xlsx')
    result['format_note']='Existing cell styles retained; advanced Excel features require manual inspection.'
    return result


def pdf_edit(ctx,a):
    policy.computer_access(ctx);policy.require(ctx,2,'Add a PDF text overlay or rotate a page; keep an original backup.')
    path,raw=fs.read_local(a['root'],a['path'])
    if path.suffix.lower()!='.pdf':raise ValueError('Choose a PDF.')
    from pypdf import PdfReader,PdfWriter
    pdf=PdfReader(BytesIO(raw))
    if pdf.is_encrypted:raise ValueError('Encrypted PDFs require manual authentication.')
    if not 1<=a['page']<=len(pdf.pages):raise ValueError('That PDF page does not exist.')
    writer=PdfWriter(clone_from=BytesIO(raw))
    page=writer.pages[a['page']-1]
    if a['operation']=='rotate':
        if a['degrees'] not in (90,180,270):raise ValueError('Use a 90, 180 or 270 degree page rotation.')
        page.rotate(a['degrees'])
    else:
        if not a['text'] or not float(page.cropbox.left)<=a['x']<float(page.cropbox.right) or not float(page.cropbox.bottom)<=a['y']<float(page.cropbox.top):raise ValueError('Put non-empty text inside the visible PDF page bounds.')
        try:a['text'].encode('latin-1')
        except UnicodeEncodeError:raise ValueError('This PDF overlay font supports Latin text only.') from None
        from reportlab.pdfgen.canvas import Canvas
        layer=BytesIO();canvas=Canvas(layer,pagesize=(float(page.mediabox.width),float(page.mediabox.height)))
        canvas.setFont('Helvetica',a['font_size']);canvas.drawString(a['x'],a['y'],a['text']);canvas.save()
        page.merge_page(PdfReader(BytesIO(layer.getvalue())).pages[0])
    out=BytesIO();writer.write(out);check=PdfReader(BytesIO(out.getvalue()))
    if len(check.pages)!=len(pdf.pages):raise OSError('PDF verification failed.')
    result=fs.save_bytes(ctx,a['root'],a['path'],out.getvalue(),a['version'],'pdf')
    result['format_note']='Text is an overlay, not replacement, redaction or reflow. Scanned text/OCR is unsupported.'
    return result


def create(ctx,a):
    policy.computer_access(ctx);policy.require(ctx,1,'Create this new ordinary document without overwriting.')
    output=BytesIO();suffix=Path(a['path']).suffix.lower()
    if suffix=='.docx':
        from docx import Document
        doc=Document()
        for line in a['content'].splitlines():doc.add_paragraph(line)
        doc.save(output)
    elif suffix=='.xlsx':
        import csv
        from openpyxl import Workbook
        book=Workbook();sheet=book.active
        for row in csv.reader(a['content'].splitlines()):
            sheet.append(["'"+v if v.startswith(('=','+','-','@')) else v for v in row])
        book.save(output);book.close()
    elif suffix=='.pdf':
        from reportlab.pdfgen.canvas import Canvas
        try:a['content'].encode('latin-1')
        except UnicodeEncodeError:raise ValueError('The basic PDF font cannot represent this text. Create DOCX for Unicode content.') from None
        canvas=Canvas(output);y=800
        for line in a['content'].splitlines():
            if y<40:canvas.showPage();y=800
            canvas.drawString(40,y,line);y-=16
        canvas.save()
    else:raise ValueError('New formatted documents support DOCX, XLSX (CSV content) or a simple Latin-text PDF.')
    return fs.save_bytes(ctx,a['root'],a['path'],output.getvalue(),'new','create_document')


def word_format(ctx,a):
    policy.computer_access(ctx);policy.require(ctx,2,'Apply this existing paragraph style to explicit paragraphs, keeping a backup.')
    path,raw=fs.read_local(a['root'],a['path'])
    if path.suffix.lower()!='.docx':raise ValueError('Paragraph formatting requires DOCX.')
    check_archive(raw)
    from docx import Document
    doc=Document(BytesIO(raw))
    if a['style'] not in doc.styles:raise ValueError('That named style is not in this document.')
    for index in a['paragraphs']:
        if not 0<=index<len(doc.paragraphs):raise ValueError('Paragraph index is outside the document body.')
        doc.paragraphs[index].style=doc.styles[a['style']]
    out=BytesIO();doc.save(out);check=Document(BytesIO(out.getvalue()))
    if any(check.paragraphs[i].style.name!=a['style'] for i in a['paragraphs']):raise OSError('Word formatting verification failed.')
    return fs.save_bytes(ctx,a['root'],a['path'],out.getvalue(),a['version'],'docx_style')
